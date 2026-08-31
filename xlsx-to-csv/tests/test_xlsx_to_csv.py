from __future__ import annotations

import csv
from datetime import datetime
from pathlib import Path

import pytest
from openpyxl import Workbook

from app.workbench import convert_workbench, find_workbooks
from app.xlsx_to_csv import XlsxToCsvError, convert_xlsx


def _save_workbook(path: Path, rows: list[list[object]]) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Data"
    for row in rows:
        worksheet.append(row)
    workbook.save(path)
    workbook.close()


def _read_csv(path: Path, encoding: str = "utf-8-sig") -> list[list[str]]:
    with path.open(encoding=encoding, newline="") as csv_file:
        return list(csv.reader(csv_file))


def test_convert_xlsx_preserves_unicode_and_csv_escaping(tmp_path: Path) -> None:
    source = tmp_path / "report.xlsx"
    _save_workbook(
        source,
        [
            ["name", "note", "active", "created"],
            ["中文", 'hello, "world"\nnext', True, datetime(2026, 8, 29, 14, 30)],
        ],
    )

    results = convert_xlsx(source)

    output = tmp_path / "report.csv"
    assert output.read_bytes().startswith(b"\xef\xbb\xbf")
    assert _read_csv(output) == [
        ["name", "note", "active", "created"],
        ["中文", 'hello, "world"\nnext', "TRUE", "2026-08-29T14:30:00"],
    ]
    assert results[0].sheet_name == "Data"
    assert results[0].row_count == 2
    assert results[0].column_count == 4


def test_default_uses_first_worksheet_and_keeps_internal_blank_rows(tmp_path: Path) -> None:
    source = tmp_path / "book.xlsx"
    workbook = Workbook()
    first = workbook.active
    first.title = "First"
    first["A1"] = "header"
    first["A3"] = "value"
    second = workbook.create_sheet("Active")
    second["A1"] = "other"
    workbook.active = 1
    workbook.save(source)
    workbook.close()

    results = convert_xlsx(source)

    assert results[0].sheet_name == "First"
    assert _read_csv(tmp_path / "book.csv") == [["header"], [], ["value"]]


def test_convert_all_sheets_uses_numbered_safe_filenames(tmp_path: Path) -> None:
    source = tmp_path / "book.xlsx"
    output = tmp_path / "csv"
    workbook = Workbook()
    first = workbook.active
    first.title = "Data"
    first["A1"] = "first"
    second = workbook.create_sheet("CON")
    second["A1"] = "second"
    second.sheet_state = "hidden"
    workbook.save(source)
    workbook.close()

    results = convert_xlsx(source, output, all_sheets=True)

    assert [result.output_path.name for result in results] == [
        "001_Data.csv",
        "002__CON.csv",
    ]
    assert _read_csv(output / "001_Data.csv") == [["first"]]
    assert _read_csv(output / "002__CON.csv") == [["second"]]


def test_select_sheet_by_name_or_one_based_index(tmp_path: Path) -> None:
    source = tmp_path / "book.xlsx"
    workbook = Workbook()
    workbook.active.title = "First"
    workbook.active["A1"] = "one"
    workbook.create_sheet("Second")["A1"] = "two"
    workbook.save(source)
    workbook.close()

    named_output = tmp_path / "named.csv"
    indexed_output = tmp_path / "indexed.csv"

    convert_xlsx(source, named_output, sheet="Second")
    convert_xlsx(source, indexed_output, sheet=2)

    assert _read_csv(named_output) == [["two"]]
    assert _read_csv(indexed_output) == [["two"]]


def test_formula_mode_exports_formula_expression(tmp_path: Path) -> None:
    source = tmp_path / "formula.xlsx"
    _save_workbook(source, [[1], ["=A1+1"]])

    convert_xlsx(source, data_only=False)

    assert _read_csv(tmp_path / "formula.csv") == [["1"], ["=A1+1"]]


def test_existing_output_requires_overwrite(tmp_path: Path) -> None:
    source = tmp_path / "report.xlsx"
    output = tmp_path / "report.csv"
    _save_workbook(source, [["new"]])
    output.write_text("old\n", encoding="utf-8")

    with pytest.raises(XlsxToCsvError, match="already exists"):
        convert_xlsx(source)
    assert output.read_text(encoding="utf-8") == "old\n"

    convert_xlsx(source, overwrite=True)
    assert _read_csv(output) == [["new"]]


def test_invalid_sheet_has_clear_error(tmp_path: Path) -> None:
    source = tmp_path / "report.xlsx"
    _save_workbook(source, [["value"]])

    with pytest.raises(XlsxToCsvError, match="was not found"):
        convert_xlsx(source, sheet="Missing")


def test_workbench_converts_every_xlsx_and_replaces_outputs(tmp_path: Path) -> None:
    workbench = tmp_path / "workbench"
    workbench.mkdir()
    first = workbench / "first.xlsx"
    second = workbench / "second.XLSX"
    workbook = Workbook()
    workbook.active.title = "Data"
    workbook.active.append(["one"])
    workbook.create_sheet("Summary").append(["summary"])
    workbook.save(first)
    workbook.close()
    _save_workbook(second, [["two"]])
    (workbench / "notes.txt").write_text("ignored", encoding="utf-8")
    (workbench / "~$locked.xlsx").write_bytes(b"ignored")

    results = convert_workbench(workbench)

    assert [path.name for path in find_workbooks(workbench)] == [
        "first.xlsx",
        "second.XLSX",
    ]
    assert len(results) == 3
    assert _read_csv(workbench / "output" / "first" / "001_Data.csv") == [["one"]]
    assert _read_csv(workbench / "output" / "first" / "002_Summary.csv") == [["summary"]]
    assert _read_csv(workbench / "output" / "second" / "001_Data.csv") == [["two"]]

    workbook = Workbook()
    workbook.active.title = "Data"
    workbook.active.append(["updated"])
    workbook.create_sheet("Summary").append(["new summary"])
    workbook.save(first)
    workbook.close()
    convert_workbench(workbench)
    assert _read_csv(workbench / "output" / "first" / "001_Data.csv") == [["updated"]]
    assert _read_csv(workbench / "output" / "first" / "002_Summary.csv") == [["new summary"]]
